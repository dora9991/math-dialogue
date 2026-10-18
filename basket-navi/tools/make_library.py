#!/usr/bin/env python3
"""
バスケ動画のライブラリ（library.js）を作る・育てる。

★ できあがる library.js には、購入した動画の名前や要約が入る。このリポジトリは公開なので、git に上げないこと
  （.gitignore で守ってある）。置き場所は、basket-navi.html と同じフォルダ（OneDrive の「バスケ動画」など）。

使い方
  # ★ いちばん簡単：走査して library.js を作り、basket-navi.html も動画フォルダに置いて、ブラウザで開く（1行）
  python3 make_library.py setup --icloud "フォルダ名" --open        # iCloud Drive（Mac）の中のフォルダ
  python3 make_library.py setup "<動画フォルダ>" --open            # それ以外（OneDrive など）
  python3 make_library.py setup --icloud                          # フォルダ名がわからないとき：動画のあるフォルダを探して並べる

  # ① フォルダを走査して、動画ぶんの項目を作る（一覧がなくてもできる）
  python3 make_library.py scan <動画フォルダ> [--base <basket-navi.html を置くフォルダ>] [--out library.js] [--probe]

  # ② Excel / CSV の一覧から作る（まとめの文章・OneDrive のリンクも取りこめる）
  python3 make_library.py list <一覧.xlsx|一覧.csv> [--sheet シート名 ...] [--map 項目=見出し,...]
                               [--videos <動画フォルダ> [--base <HTMLの置き場所>]] [--out library.js]

  --probe   ffprobe があれば、動画の長さ（秒）も入れる（iCloud で Mac にまだ無い動画は、ダウンロードが始まらないよう飛ばす）
  --open    できあがった basket-navi.html をブラウザで開く（Mac は Chrome → Edge → 既定の順）
  --map     見出しの自動判定が合わないとき。項目は title / category / file / url / summary / tags / source
            例：--map title=動画名,summary=内容
  --videos  一覧にファイルの場所がないとき、動画フォルダから「名前が一致する動画」を探して file を埋める

  すでに library.js があれば、手で書いた要約・見どころ・タグを消さずに「足す」だけ。
  上書きの前に library.js.bak に控えを作る。
"""

import argparse
import csv
import json
import os
import re
import shutil
import subprocess
import sys
import unicodedata
from pathlib import Path

ICLOUD_ROOT = Path.home() / 'Library' / 'Mobile Documents' / 'com~apple~CloudDocs'   # Mac の iCloud Drive の入り口
SF_DATALESS = 0x40000000   # macOS：iCloud 上だけにあって、まだ Mac にダウンロードされていないファイルの印
VIDEO_EXT = {'.mp4', '.m4v', '.mov', '.webm', '.mkv', '.avi', '.wmv', '.mpg', '.mpeg', '.flv', '.ts', '.m2ts'}
CATS = [{'id': 'coaching', 'name': '指導法'}, {'id': 'tactics', 'name': '戦術'}, {'id': 'training', 'name': 'トレーニング'}]
CAT_NAMES = {c['name']: c['id'] for c in CATS}

# 名前からの分類の当たりづけ。「まちがえて入れる」より「未分類にしておく」を選ぶ（同点・該当なしは空）
CAT_WORDS = {
    'coaching': ['指導', 'コーチ', 'マネジメント', 'チームづくり', 'チーム作り', '育成', '声かけ', '声掛け', 'メンタル', '運営', '練習計画', '考え方'],
    'tactics': ['戦術', 'オフェンス', 'ディフェンス', 'セット', 'ゾーン', 'ピック', 'トランジション', 'プレー', 'フォーメーション', '速攻', 'スペーシング', 'ポジション', 'マンツーマン', 'プレス'],
    'training': ['トレーニング', 'ドリル', 'フィジカル', 'ハンドリング', 'シュート', 'ステップ', '体力', 'ウォームアップ', 'ストレッチ', 'メニュー', '基礎', '筋', 'コーディネーション'],
}

# 一覧（Excel / CSV）の見出しの呼び方。1つ目の判定は「完全一致」、2つ目は「ふくむ」
HEADERS = {
    'title': ['タイトル', '題名', '動画名', '動画タイトル', '名称', '名前', '動画', 'ファイル名', 'title', 'name'],
    'category': ['カテゴリ', 'カテゴリー', '分類', 'ジャンル', '種別', '区分', '種類', 'category'],
    'file': ['ファイル', 'ファイルパス', 'パス', '保存先', '場所', 'file', 'path'],
    'url': ['url', 'リンク', '共有リンク', 'onedrive', '動画リンク'],
    'summary': ['概要', '要約', 'まとめ', '内容', '説明', '要点', '備考', 'メモ', 'summary'],
    'tags': ['タグ', 'キーワード', 'テーマ', 'tags'],
    'source': ['講師', '出典', '購入元', '提供元', 'シリーズ', '提供', 'source'],
}
FIELD_ORDER = ['title', 'category', 'file', 'url', 'summary', 'tags', 'source']

HEADER_COMMENT = """/* バスケ指導ナビのライブラリ（自分用・非公開）。★git に上げない（.gitignore 済み）。
   手で直してよい。tools/make_library.py を再実行しても、ここに書いた要約・見どころ・タグは消えない。
   書き方は data/sample.js の先頭のコメントを見る。 */
"""


def die(msg):
    print('エラー：' + msg, file=sys.stderr)
    sys.exit(1)


def nfc(s):
    return unicodedata.normalize('NFC', s)


def key(s):
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', str(s))).lower()


def natural_key(p):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r'(\d+)', unicodedata.normalize('NFKC', str(p)))]


# ---------- library.js の読み書き ----------
def read_lib(path):
    if not path.exists():
        return None
    txt = path.read_text(encoding='utf-8')
    a, b = txt.find('{'), txt.rfind('}')
    if a < 0 or b < a:
        die(f'既存の {path} を読めません（形がちがう）。上書きしないで止めました。')
    try:
        lib = json.loads(txt[a:b + 1])
    except Exception as e:  # noqa: BLE001
        die(f'既存の {path} が JSON として読めません（{e}）。手で書き換えたなら、カンマ・引用符を確認してください。上書きしないで止めました。')
    if not isinstance(lib.get('videos'), list):
        die(f'既存の {path} に videos がありません。上書きしないで止めました。')
    return lib


def write_lib(path, lib):
    if path.exists():
        shutil.copy2(path, str(path) + '.bak')
    path.parent.mkdir(parents=True, exist_ok=True)
    body = json.dumps(lib, ensure_ascii=False, indent=1)
    path.write_text(HEADER_COMMENT + 'window.BASKET_LIBRARY = ' + body + ';\n', encoding='utf-8')


def new_lib():
    return {'version': 1, 'title': 'マイ・バスケ動画', 'categories': CATS, 'videos': []}


def next_id_number(videos):
    nums = [int(m.group(1)) for v in videos for m in [re.fullmatch(r'v(\d+)', str(v.get('id', '')))] if m]
    return (max(nums) + 1) if nums else 1


def new_video(n, title, category, file, tags=None, summary='', source='', url='', duration=0):
    return {'id': f'v{n:03d}', 'title': title, 'category': category, 'file': file, 'duration': duration,
            'tags': tags or [], 'summary': summary, 'source': source, 'url': url, 'clips': []}


# ---------- 分類 ----------
def cat_score(text, weight):
    t = unicodedata.normalize('NFKC', text)
    return {cid: weight * sum(1 for w in words if w in t) for cid, words in CAT_WORDS.items()}


def guess_category(parts, title):
    score = {cid: 0 for cid in CAT_WORDS}
    for part in parts:
        if part in CAT_NAMES:                       # フォルダ名が「戦術」のようにそのもの
            score[CAT_NAMES[part]] += 5
        for cid, s in cat_score(part, 3).items():
            score[cid] += s
    for cid, s in cat_score(title, 1).items():
        score[cid] += s
    best = max(score.values())
    winners = [cid for cid, s in score.items() if s == best]
    return winners[0] if best > 0 and len(winners) == 1 else ''


def cat_from_text(s):
    s = nfc(str(s)).strip()
    if not s:
        return ''
    if s in CAT_NAMES:
        return CAT_NAMES[s]
    if s in {c['id'] for c in CATS}:
        return s
    return guess_category([s], '')


# ---------- 動画の長さ ----------
_probe_hint = [False]


def probe(path):
    try:
        r = subprocess.run(['ffprobe', '-v', 'error', '-show_entries', 'format=duration', '-of', 'default=nw=1:nk=1', str(path)],
                           capture_output=True, text=True, timeout=60)
        return round(float(r.stdout.strip())) if r.returncode == 0 and r.stdout.strip() else 0
    except FileNotFoundError:
        if not _probe_hint[0]:
            print('メモ：ffprobe が見つからないので、動画の長さは入れていません（ffmpeg を入れると入ります）。')
            _probe_hint[0] = True
        return 0
    except Exception:  # noqa: BLE001
        return 0


# ---------- ① scan ----------
def is_dataless(p):
    """iCloud Drive で、Mac にまだダウンロードされていないファイルか（macOS だけ。ほかの OS では常に False）。"""
    try:
        return bool(getattr(os.stat(p), 'st_flags', 0) & SF_DATALESS)
    except OSError:
        return False


def icloud_placeholder_name(fn):
    """古い macOS の iCloud Drive は、未ダウンロードの動画を「.名前.icloud」という小さな目印のファイルにする。本来の名前を返す（目印でなければ空）。"""
    return fn[1:-len('.icloud')] if (fn.startswith('.') and fn.endswith('.icloud') and len(fn) > len('..icloud')) else ''


def walk_videos(root):
    """(動画のパス, まだ Mac にダウンロードされていないか) を返す。目印のファイルも動画として数える。"""
    for dp, dns, fns in os.walk(root):
        dns[:] = [d for d in dns if not d.startswith('.') and d not in ('@eaDir', '__MACOSX')]
        for fn in fns:
            if fn.startswith('._'):
                continue
            real = icloud_placeholder_name(fn)
            if real:
                if os.path.splitext(real)[1].lower() in VIDEO_EXT:
                    yield Path(dp) / real, True
                continue
            if fn.startswith('.'):
                continue
            if os.path.splitext(fn)[1].lower() in VIDEO_EXT:
                p = Path(dp) / fn
                yield p, is_dataless(p)


def present(p):
    """実ファイルがあるか、iCloud の未ダウンロードの目印があるか。"""
    return p.exists() or (p.parent / ('.' + p.name + '.icloud')).exists()


def tags_from(parts):
    out = []
    for p in parts:
        p = nfc(p).strip()
        if p and p not in CAT_NAMES and p not in out:
            out.append(p)
    return out[:3]


def cmd_scan(a):
    root = Path(a.folder).expanduser().resolve()
    if not root.is_dir():
        die(f'フォルダが見つかりません：{root}')
    base = Path(a.base).expanduser().resolve() if a.base else root
    try:
        root.relative_to(base)
    except ValueError:
        die('--base は、動画フォルダと同じか、その「上」のフォルダにしてください（動画が HTML より下にないと、ブラウザで再生できません）')
    out = Path(a.out).expanduser() if a.out else base / 'library.js'
    lib = read_lib(out) or new_lib()
    have = {key(v['file']): v for v in lib['videos'] if v.get('file')}
    n = next_id_number(lib['videos'])
    found = added = offline_n = 0
    for p, offline in sorted(walk_videos(root), key=lambda t: natural_key(t[0])):
        found += 1
        offline_n += 1 if offline else 0
        rel = p.relative_to(base).as_posix()
        v = have.get(key(rel))
        can_probe = a.probe and not offline          # iCloud の未ダウンロードを probe すると、ダウンロードが始まってしまう
        if v is not None:
            if can_probe and not v.get('duration'):
                v['duration'] = probe(p)
            continue
        parts = list(p.relative_to(root).parts[:-1])
        title = nfc(p.stem)
        lib['videos'].append(new_video(n, title, guess_category(parts, title), rel,
                                       tags=tags_from(parts), duration=probe(p) if can_probe else 0))
        n += 1
        added += 1
    missing = [v for v in lib['videos'] if v.get('file') and not present(base / v['file'])]
    write_lib(out, lib)
    print(f'動画 {found} 本を見つけた ／ 新しく追加 {added} 本 ／ もとから入っていた {found - added} 本')
    if offline_n:
        print(f'iCloud Drive で、まだ Mac にダウンロードされていない動画 {offline_n} 本：Finder で動画（またはフォルダ）を右クリック →「今すぐダウンロード」にすると、再生できます。'
              + ('動画の長さは、ダウンロード後に --probe 付きでもう一度実行すると入ります。' if a.probe else ''))
    unclassified = sum(1 for v in lib['videos'] if not v.get('category'))
    if unclassified:
        print(f'分類がわからなかった動画 {unclassified} 本 → 「未分類」で入っています。library.js の category を直すと分類されます')
    if missing:
        print(f'ファイルが見つからない項目 {len(missing)} 件（移動・名前変更の可能性）：' + '、'.join(v['title'] for v in missing[:5]) + (' …' if len(missing) > 5 else ''))
    print(f'書き出した：{out}')
    if not getattr(a, 'from_setup', False):
        print('basket-navi.html を同じフォルダに置いて、開いてください。')


# ---------- ② list ----------
def cell_val(c):
    v = getattr(c, 'value', c)
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def cell_link(c, formula_cell=None):
    h = getattr(c, 'hyperlink', None)
    t = getattr(h, 'target', None) if h else None
    if t and re.match(r'https?://', t):
        return t
    f = getattr(formula_cell, 'value', None)
    if isinstance(f, str) and f.upper().startswith('=HYPERLINK('):
        m = re.search(r'"(https?://[^"]+)"', f)
        if m:
            return m.group(1)
    return ''


def norm_head(s):
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', str(s))).lower()


def assign_columns(heads, mapping):
    cols, used = {}, set()
    for field, head in mapping.items():          # --map の指定がいちばん強い
        for i, h in enumerate(heads):
            if norm_head(h) == norm_head(head):
                cols[field] = i
                used.add(i)
    for exact in (True, False):                   # 完全一致 → ふくむ
        for field in FIELD_ORDER:
            if field in cols:
                continue
            for i, h in enumerate(heads):
                if i in used or not h:
                    continue
                nh = norm_head(h)
                if any((nh == norm_head(w)) if exact else (norm_head(w) in nh) for w in HEADERS[field]):
                    cols[field] = i
                    used.add(i)
                    break
    return cols


def table_rows(sheet, grid, header_row, mapping, formulas=None):
    """grid は 行×セル。セルは openpyxl のセルか文字列。"""
    hi = header_row - 1 if header_row else next((i for i, r in enumerate(grid) if sum(1 for c in r if cell_val(c)) >= 2), None)
    if hi is None or hi >= len(grid):
        return
    heads = [cell_val(c) for c in grid[hi]]
    cols = assign_columns(heads, mapping)
    if 'title' not in cols:
        die(f'シート「{sheet}」で、動画のタイトルの列がわかりません。見出し：{heads}\n  → --map title=（タイトルにあたる見出し） を付けて、もう一度。')
    for ri in range(hi + 1, len(grid)):
        row = grid[ri]

        def get(field):
            i = cols.get(field)
            return cell_val(row[i]) if i is not None and i < len(row) else ''
        title = get('title')
        if not title:
            continue
        link = ''
        for i in ([cols['url']] if 'url' in cols else []) + [cols['title']]:
            if i < len(row):
                fc = formulas[ri][i] if formulas and ri < len(formulas) and i < len(formulas[ri]) else None
                link = cell_link(row[i], fc) or link
        url = get('url')
        yield {'sheet': sheet, 'title': title, 'category': get('category'), 'file': get('file'),
               'url': url if re.match(r'https?://', url) else link, 'summary': get('summary'),
               'tags': get('tags'), 'source': get('source')}


def read_list(path, sheets, header_row, mapping):
    ext = path.suffix.lower()
    if ext in ('.xlsx', '.xlsm'):
        try:
            import openpyxl
        except ImportError:
            die('openpyxl が必要です：pip3 install openpyxl')
        wb = openpyxl.load_workbook(path, data_only=True)
        wf = openpyxl.load_workbook(path)
        rows = []
        for ws in wb.worksheets:
            if sheets and ws.title not in sheets:
                continue
            grid = [list(r) for r in ws.iter_rows()]
            formulas = [list(r) for r in wf[ws.title].iter_rows()]
            rows += list(table_rows(ws.title, grid, header_row, mapping, formulas))
        return rows
    if ext in ('.csv', '.tsv', '.txt'):
        raw = path.read_bytes()
        for enc in ('utf-8-sig', 'cp932'):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            die('CSV の文字コードを判定できません（UTF-8 か Shift_JIS にしてください）')
        delim = '\t' if ext in ('.tsv', '.txt') else ','
        grid = [r for r in csv.reader(text.splitlines(), delimiter=delim)]
        return list(table_rows(path.stem, grid, header_row, mapping))
    die('一覧は .xlsx / .csv で渡してください（Word や PDF の一覧は、Excel か CSV に貼りなおすのがいちばん確実です）')


def find_file_for(title, files, root):
    """タイトルと「ほぼ同じ名前」の動画を1つだけ探す。複数・なしなら空。"""
    kt = key(title)
    exact = [f for f in files if key(f.stem) == kt]
    near = exact or [f for f in files if kt and (kt in key(f.stem) or key(f.stem) in kt) and len(key(f.stem)) >= 4]
    return near[0] if len(near) == 1 else None


def cmd_list(a):
    src = Path(a.list_file).expanduser()
    if not src.exists():
        die(f'一覧が見つかりません：{src}')
    mapping = {}
    for pair in filter(None, (a.map or '').split(',')):
        if '=' not in pair or pair.split('=')[0].strip() not in HEADERS:
            die(f'--map は「項目=見出し」の形で。項目は {", ".join(HEADERS)}（例：--map title=動画名,summary=内容）')
        f, h = pair.split('=', 1)
        mapping[f.strip()] = h.strip()
    rows = read_list(src, set(a.sheet or []), a.header_row, mapping)
    if not rows:
        die('一覧から、動画の行を読めませんでした。')
    root = Path(a.videos).expanduser().resolve() if a.videos else None
    base = Path(a.base).expanduser().resolve() if a.base else root
    files = sorted((p for p, _ in walk_videos(root)), key=natural_key) if root else []
    out = Path(a.out).expanduser() if a.out else (base / 'library.js' if base else Path('library.js'))
    lib = read_lib(out) or new_lib()
    by_title = {key(v['title']): v for v in lib['videos']}
    by_file = {key(v['file']): v for v in lib['videos'] if v.get('file')}
    n = next_id_number(lib['videos'])
    added = merged = unlinked = 0
    for r in rows:
        file = r['file'].replace('\\', '/').strip()
        if not file and files:
            f = find_file_for(r['title'], files, root)
            if f is not None:
                file = f.relative_to(base).as_posix()
        category = cat_from_text(r['category']) or cat_from_text(r['sheet']) or guess_category([], r['title'])
        tags = [t.strip() for t in re.split(r'[,、，/／]', r['tags']) if t.strip()]
        ctext = nfc(r['category']).strip()
        if ctext and ctext not in CAT_NAMES and ctext not in {c['id'] for c in CATS} and ctext not in tags:
            tags.append(ctext)                       # 「戦術」などそのものでない分類の語は、タグとしても残す
        v = (by_file.get(key(file)) if file else None) or by_title.get(key(r['title']))
        if v is not None:                              # もとの項目は壊さず、空いているところだけ埋める
            for f, val in (('summary', r['summary']), ('source', r['source']), ('url', r['url']), ('file', file)):
                if val and not v.get(f):
                    v[f] = val
            if not v.get('category') and category:
                v['category'] = category
            v['tags'] = (v.get('tags') or []) + [t for t in tags if t not in (v.get('tags') or [])]
            merged += 1
            continue
        lib['videos'].append(new_video(n, nfc(r['title']), category, file, tags=tags, summary=r['summary'], source=r['source'], url=r['url']))
        n += 1
        added += 1
        if not file:
            unlinked += 1
    write_lib(out, lib)
    print(f'一覧の {len(rows)} 行を読んだ ／ 新しく追加 {added} 本 ／ もとの項目に反映 {merged} 本')
    if unlinked:
        print(f'動画ファイルにひも付けられなかった項目 {unlinked} 件。file を書くか、--videos <動画フォルダ> を付けて再実行すると、名前が合うものを探します')
    print(f'書き出した：{out}')


# ---------- ★ setup ----------
def count_videos(dirpath, cap_files=30000):
    """動画の数（iCloud の未ダウンロードの目印も数える）。大きなフォルダでは途中でやめる。"""
    n = seen = 0
    for dp, dns, fns in os.walk(dirpath):
        dns[:] = [d for d in dns if not d.startswith('.')]
        for fn in fns:
            seen += 1
            name = icloud_placeholder_name(fn) or fn
            if not name.startswith('.') and os.path.splitext(name)[1].lower() in VIDEO_EXT:
                n += 1
            if seen >= cap_files:
                return n
    return n


def pick_icloud_folder(root, name):
    if not root.is_dir():
        die(f'iCloud Drive の場所が見つかりません：{root}\n  → Mac で iCloud Drive をオンにして、動画をそこに入れてください。Mac 以外では、動画のフォルダを直接指定してください。')
    dirs = sorted((d for d in root.iterdir() if d.is_dir() and not d.name.startswith('.')), key=lambda d: natural_key(d.name))
    if name:
        hit = next((d for d in dirs if key(d.name) == key(name)), None)
        if hit is None:
            die(f'iCloud Drive に「{name}」というフォルダがありません。いまあるフォルダ：' + '、'.join(nfc(d.name) for d in dirs[:30]))
        return hit
    cands = sorted(((count_videos(d), d) for d in dirs), key=lambda t: -t[0])
    cands = [(n, d) for n, d in cands if n > 0]
    if not cands:
        die('iCloud Drive の中に、動画のあるフォルダが見つかりませんでした。動画を入れたフォルダの名前を、--icloud の後ろに書いてください。')
    if len(cands) > 1:
        print('iCloud Drive に、動画のあるフォルダが複数あります。使うフォルダの名前を指定して、もう一度実行してください：')
        for n, d in cands[:10]:
            print(f'  {n:>4} 本  {nfc(d.name)}')
        print(f'  例：python3 tools/make_library.py setup --icloud "{nfc(cands[0][1].name)}" --open')
        sys.exit(2)
    print(f'iCloud Drive の「{nfc(cands[0][1].name)}」（動画 {cands[0][0]} 本）を使います。')
    return cands[0][1]


def open_in_browser(path):
    """できあがった HTML をブラウザで開く。開いたブラウザの名前を返す。"""
    if sys.platform == 'darwin':
        for app in ('Google Chrome', 'Microsoft Edge'):
            if subprocess.run(['open', '-a', app, str(path)], capture_output=True).returncode == 0:
                return app
        subprocess.run(['open', str(path)], check=True)
    elif sys.platform.startswith('win'):
        os.startfile(str(path))  # noqa: S606
    else:
        subprocess.run(['xdg-open', str(path)], check=True)
    return '既定のブラウザ'


def cmd_setup(a):
    root = Path(a.icloud_root).expanduser() if a.icloud_root else ICLOUD_ROOT
    if a.icloud is not None:
        if a.folder:
            die('--icloud と、フォルダの指定は、どちらか片方にしてください。')
        folder = pick_icloud_folder(root, a.icloud)
    elif a.folder:
        folder = Path(a.folder).expanduser()
    else:
        die('動画のフォルダを指定してください（iCloud Drive なら --icloud ["フォルダ名"]）。\n  例：python3 tools/make_library.py setup --icloud "バスケ動画" --open')
    folder = folder.resolve()
    if not folder.is_dir():
        die(f'フォルダが見つかりません：{folder}')
    src = Path(__file__).resolve().parent.parent / 'basket-navi.html'
    if not src.exists():
        die(f'{src} がありません。先に node tools/bundle.cjs を実行してください。')
    cmd_scan(argparse.Namespace(folder=str(folder), base=None, out=None, probe=a.probe, from_setup=True))
    dst = folder / 'basket-navi.html'
    same = dst.exists() and dst.read_bytes() == src.read_bytes()
    if not same:
        shutil.copy2(src, dst)
    print(f'basket-navi.html は{"すでに最新です" if same else "動画フォルダに置きました"}：{dst}')
    if a.open:
        try:
            print(f'{open_in_browser(dst)} で開きました。')
        except Exception as e:  # noqa: BLE001
            print(f'ブラウザを開けませんでした（{e}）。{dst} をダブルクリックしてください。')
    else:
        print('basket-navi.html をダブルクリックして開いてください（Chrome か Edge がおすすめ）。')


def main():
    ap = argparse.ArgumentParser(description='バスケ動画のライブラリ（library.js）を作る・育てる', formatter_class=argparse.RawDescriptionHelpFormatter, epilog=__doc__)
    sub = ap.add_subparsers(dest='cmd', required=True)
    su = sub.add_parser('setup', help='走査して library.js を作り、basket-navi.html も動画フォルダに置いて開く（いちばん簡単）')
    su.add_argument('folder', nargs='?')
    su.add_argument('--icloud', nargs='?', const='', default=None, metavar='フォルダ名', help='iCloud Drive（Mac）の中のフォルダ。名前を省くと、動画のあるフォルダを探す')
    su.add_argument('--icloud-root', help=argparse.SUPPRESS)
    su.add_argument('--probe', action='store_true')
    su.add_argument('--open', action='store_true')
    su.set_defaults(fn=cmd_setup)
    s = sub.add_parser('scan', help='動画フォルダを走査して項目を作る')
    s.add_argument('folder')
    s.add_argument('--base')
    s.add_argument('--out')
    s.add_argument('--probe', action='store_true')
    s.set_defaults(fn=cmd_scan)
    li = sub.add_parser('list', help='Excel / CSV の一覧から作る')
    li.add_argument('list_file')
    li.add_argument('--sheet', action='append')
    li.add_argument('--header-row', type=int, default=0, help='見出し行の番号（1始まり。省略すると自動）')
    li.add_argument('--map')
    li.add_argument('--videos')
    li.add_argument('--base')
    li.add_argument('--out')
    li.set_defaults(fn=cmd_list)
    a = ap.parse_args()
    a.fn(a)


if __name__ == '__main__':
    main()
